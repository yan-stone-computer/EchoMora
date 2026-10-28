from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "docs" / "architecture.zh-CN.md"
OUTPUT = ROOT / "docs" / "EchoMora架构详细说明.docx"
ASSETS = ROOT / "docs" / "assets"
TOOLS = Path(
    r"C:\Users\34274\.codex\plugins\cache\openai-primary-runtime"
    r"\documents\26.805.11740\skills\documents\scripts"
)
sys.path.insert(0, str(TOOLS))

from table_geometry import (  # noqa: E402
    apply_table_geometry,
    column_widths_from_weights,
    section_content_width_dxa,
)


INK = RGBColor(0x18, 0x23, 0x26)
TEAL = RGBColor(0x1F, 0x6B, 0x66)
TEAL_DARK = RGBColor(0x18, 0x55, 0x51)
AMBER = RGBColor(0x9B, 0x67, 0x28)
MUTED = RGBColor(0x63, 0x72, 0x74)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE_TEAL = "E8F1EF"
PALE_AMBER = "F7ECDA"
PALE_SLATE = "F0F3F5"
CODE_FILL = "F3F5F4"


def set_run_font(run, name: str, size: float | None = None, color=None, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def shade_paragraph(paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    p_pr.append(shd)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("EchoMora  ·  ")
    set_run_font(run, "Microsoft YaHei", 8.5, MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_picture_alt(paragraph, filename: str, description: str, width_mm: float):
    run = paragraph.add_run()
    run.add_picture(str(ASSETS / filename), width=Mm(width_mm))
    for doc_pr in run._r.xpath(".//wp:docPr"):
        doc_pr.set("descr", description)
        doc_pr.set("title", filename)


def add_inline_runs(paragraph, text: str, base_size: float = 10.5) -> None:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace(r"\(", "").replace(r"\)", "")
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])
            set_run_font(run, "Microsoft YaHei", base_size, INK)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, "Microsoft YaHei", base_size, INK, True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Cascadia Mono", base_size - 0.8, TEAL_DARK)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, "Microsoft YaHei", base_size, INK)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, TEAL, 18, 8),
        "Heading 2": (12.5, TEAL_DARK, 12, 5),
        "Heading 3": (11, AMBER, 10, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    caption.base_style = normal
    caption.font.size = Pt(9)
    caption.font.color.rgb = MUTED
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.keep_with_next = False

    code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = normal
    code.font.name = "Cascadia Mono"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(8.7)
    code.font.color.rgb = INK
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.08
    code.paragraph_format.left_indent = Mm(3)
    code.paragraph_format.right_indent = Mm(3)


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    add_picture_alt(
        p,
        "cover.png",
        "群聊、神经关系网络与情绪信号的抽象概念图，仅作封面视觉，不表达系统连线。",
        168,
    )

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(4)
    run = title.add_run("EchoMora")
    set_run_font(run, "Bahnschrift", 30, INK, True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    run = subtitle.add_run("分层记忆与情感协同架构详设")
    set_run_font(run, "Microsoft YaHei", 19, TEAL_DARK, True)

    quote = doc.add_paragraph()
    quote.alignment = WD_ALIGN_PARAGRAPH.CENTER
    quote.paragraph_format.space_after = Pt(22)
    run = quote.add_run("不是聊天记录，是会靠近、也会淡忘的 AI 群友。")
    set_run_font(run, "Microsoft YaHei", 11.5, MUTED)

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rule.paragraph_format.space_after = Pt(18)
    run = rule.add_run("MEMORY  ·  GRAPHRAG  ·  AFFECT  ·  SOFT FORGETTING")
    set_run_font(run, "Bahnschrift", 9, AMBER, True)

    rows = [
        ("文档版本", "v0.1.0"),
        ("适用场景", "QQ / 微信群机器人 · 项目答辩 · 技术方案评审"),
        ("更新日期", "2026-10-28"),
    ]
    for label, value in rows:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(4)
        label_run = paragraph.add_run(f"{label}  ")
        set_run_font(label_run, "Microsoft YaHei", 9.3, TEAL_DARK, True)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, "Microsoft YaHei", 9.3, INK)

    doc.add_page_break()


def add_table(doc: Document, lines: list[str]) -> None:
    rows = [
        [cell.strip() for cell in line.strip().strip("|").split("|")]
        for line in lines
        if not re.match(r"^\s*\|?[\s:|-]+\|?\s*$", line)
    ]
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for col_index in range(cols):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(3)
            text = row[col_index] if col_index < len(row) else ""
            add_inline_runs(paragraph, text, 9.2)
            if row_index == 0:
                set_cell_shading(cell, PALE_TEAL)
                for run in paragraph.runs:
                    run.bold = True
    header_properties = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_properties.append(repeat_header)
    width = section_content_width_dxa(doc.sections[-1])
    weights = [1.2 if index == 0 else 1.0 for index in range(cols)]
    apply_table_geometry(
        table,
        column_widths_from_weights(weights, width),
        table_width_dxa=width,
    )
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(3)


def create_number_definition(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    numeric_ids = [
        int(item.get(qn("w:numId")))
        for item in numbering.findall(qn("w:num"))
        if item.get(qn("w:numId"))
    ]
    num_id = max(numeric_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), "7")
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(abstract)
    num.append(override)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.append(level)
    num_pr.append(number)
    p_pr.append(num_pr)


def add_code(doc: Document, code_lines: list[str]) -> None:
    if not code_lines:
        return
    for index, line in enumerate(code_lines):
        paragraph = doc.add_paragraph(style="Code Block")
        shade_paragraph(paragraph, CODE_FILL)
        run = paragraph.add_run(line or " ")
        set_run_font(run, "Cascadia Mono", 8.7, INK)
        if index == 0:
            paragraph.paragraph_format.space_before = Pt(5)
        if index == len(code_lines) - 1:
            paragraph.paragraph_format.space_after = Pt(8)


def add_body(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    skipped_first_title = False
    pending_caption = False
    numbered_list_active = False
    current_num_id: int | None = None
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()

        if not stripped:
            numbered_list_active = False
            index += 1
            continue

        if stripped.startswith("```"):
            numbered_list_active = False
            fence = stripped[3:].strip().lower()
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            index += 1
            if fence != "mermaid":
                add_code(doc, code_lines)
            continue

        image = re.match(r"!\[[^\]]*\]\(assets/([^)]+)\)", stripped)
        if image:
            numbered_list_active = False
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(4)
            add_picture_alt(
                paragraph,
                image.group(1),
                "EchoMora 分层记忆与情感协同架构图。",
                168,
            )
            pending_caption = True
            index += 1
            continue

        if stripped.startswith("#"):
            numbered_list_active = False
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and not skipped_first_title:
                skipped_first_title = True
                index += 1
                while index < len(lines) and (
                    not lines[index].strip() or lines[index].strip().startswith(">")
                ):
                    index += 1
                continue
            doc.add_heading(text, level=min(3, max(1, level)))
            index += 1
            continue

        if stripped.startswith("|"):
            numbered_list_active = False
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, table_lines)
            continue

        if pending_caption and stripped.startswith("图 "):
            paragraph = doc.add_paragraph(style="Figure Caption")
            add_inline_runs(paragraph, stripped, 9)
            pending_caption = False
            index += 1
            continue

        if stripped.startswith("\\["):
            numbered_list_active = False
            math_lines = [stripped]
            while not math_lines[-1].endswith("\\]") and index + 1 < len(lines):
                index += 1
                math_lines.append(lines[index].strip())
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            text = " ".join(math_lines).replace("\\[", "").replace("\\]", "")
            run = paragraph.add_run(text)
            set_run_font(run, "Cambria Math", 10.5, TEAL_DARK)
            index += 1
            continue

        if stripped.startswith("> "):
            index += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            numbered_list_active = False
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline_runs(paragraph, re.sub(r"^[-*]\s+", "", stripped))
            index += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            if not numbered_list_active:
                current_num_id = create_number_definition(doc)
                numbered_list_active = True
            paragraph = doc.add_paragraph(style="List Number")
            if current_num_id is not None:
                apply_numbering(paragraph, current_num_id)
            add_inline_runs(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
            index += 1
            continue

        numbered_list_active = False
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith("#")
                or candidate.startswith("|")
                or candidate.startswith("```")
                or candidate.startswith("![")
                or candidate.startswith(">")
                or candidate.startswith("\\[")
                or re.match(r"^[-*]\s+", candidate)
                or re.match(r"^\d+\.\s+", candidate)
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_inline_runs(paragraph, " ".join(paragraph_lines))


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(10)

    configure_styles(doc)
    add_page_number(section.footer.paragraphs[0])
    add_cover(doc)
    add_body(doc, SOURCE.read_text(encoding="utf-8"))

    doc.core_properties.title = "EchoMora 分层记忆与情感协同架构"
    doc.core_properties.subject = "群聊 AI 记忆、GraphRAG、情感状态与软遗忘技术方案"
    doc.core_properties.author = "EchoMora contributors"
    doc.core_properties.keywords = "memory, GraphRAG, DeepSeek, chatbot, soft forgetting"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()

